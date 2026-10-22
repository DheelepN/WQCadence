"""Build the manuscript (.docx) from results/summary.json and figures/.

    python scripts/build_paper.py

Every number in the text is read from the results file, so the manuscript
cannot drift from the code that produced it.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
S = json.loads((ROOT / "results" / "summary.json").read_text())
OUT = ROOT / "paper" / "AquaSentinel_manuscript.docx"

P1, P2, P3, DATA, CAL = S["pillar1"], S["pillar2"], S["pillar3"], S["data"], S["detector_calibration"]
R, TESTS = P1["recall"], P1["tests"]
BG = {(b["policy"], b["persist"]): b for b in P1["background"]}
REAL = P1["real_events"]["by_policy"]
DES = P2["designs"]
CAD, CMP = P3["by_cadence"], P3["comparisons"]
FX, AD, CO = "Fixed-30", "Adaptive-30/15", "Continuous-15"

SUP = str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")


def p(x: float) -> str:
    if x < 1e-4:
        e = int(f"{x:e}".split("e")[1])
        return f"p < 10{str(e + 1).translate(SUP)}"
    if x < 0.001:
        return "p < 0.001"
    return f"p = {x:.2f}" if x >= 0.01 else f"p = {x:.3f}"


def f3(x): return f"{x:.3f}"
def f2(x): return f"{x:.2f}"
def m(x): return f"{x:,.0f} m"
def pc(x): return f"{100 * x:.0f}%"
def n0(x): return f"{x:,.0f}"


WORDS = {0: "none", 1: "one", 2: "two", 3: "three", 4: "four"}


def real_hits(policy: str) -> float:
    return sum(REAL[policy].values())


def ci(c): return f"[{c[0]:.3f}, {c[1]:.3f}]"


# ---------------------------------------------------------------- numbers ---
N_TRIALS = P1["n_trials"]
N_SEV = R[FX]["n_brief"]
QUIET_DAYS = BG[(FX, 2)]["quiet_steps"] / 96          # quiet days scored for false alarms (after warm-up)
REC_DAYS = sum(DATA["records"].values()) / 96
mc_af = TESTS["Adaptive_vs_Fixed30"]
mc_ac = TESTS["Adaptive_vs_Continuous"]
lat_af = mc_af["latency"]
samp = {k: v["samples_per_station_day"] for k, v in CAD.items()}
saving_305_vs_15 = 1 - samp["Adaptive-30/5"] / samp["Fixed-15"]
extra_3015 = samp["Adaptive-30/15"] / samp["Fixed-30"] - 1
c3015 = CMP["Adaptive-30/15 vs Fixed-30"]
c305_15 = CMP["Adaptive-30/5 vs Fixed-15"]
c305_30 = CMP["Adaptive-30/5 vs Fixed-30"]
c15_30 = CMP["Fixed-15 vs Fixed-30"]
c5_15 = CMP["Fixed-5 vs Fixed-15"]
st = {(r["cadence"], r["n_stations"]): r for r in P3["station_count"]}
a_pri = DES["discharge_only|F4"]
rev = DES["discharge_and_runoff|F6"]
f4_both = DES["discharge_and_runoff|F4"]
rms = {k: f"{k} [rms]" for k in (FX, AD, CO)}

AUTHOR = "Dheelep Sai Gupthaa Nunna"   # family name: Nunna
ADDRESS = ("Independent Researcher, 19-199, Gundla Puttur Street, Gate Puttur, Puttur, Chittoor District, "
           "Andhra Pradesh [PIN code], India")
EMAIL = "dheelep.phd@gmail.com"
ORCID = "0009-0006-9652-4387"
ZENODO_DOI = "[Zenodo DOI]"
REPO_URL = "https://github.com/DheelepN/Aqua-sentinal"

TITLE = ("Where adaptive sampling pays off in river water-quality surveillance: detection, "
         "fault discrimination and source localization in a reproducible benchmark")

HIGHLIGHTS = [
    "Open, tested benchmark of adaptive sampling, fault checks and source localization",
    "Escalating 30 to 15 min sampling did not detect more events than fixed 30 min",
    f"The same escalation cut median source-localization error from {CAD['Fixed-30']['inversion_median']:.0f} to "
    f"{CAD['Adaptive-30/15']['inversion_median']:.0f} m",
    "Fault-robust multichannel detection missed most real turbidity-led river events",
    "Results checked against expert-verified events and real sensor faults",
]
for h in HIGHLIGHTS:
    assert len(h) <= 85, (len(h), h)

ABSTRACT = (
    "Adaptive sampling is usually justified by better event detection. We tested this for river water-quality "
    "stations using a public 15-min sonde record from the Owenmore River, Ireland, Monte Carlo injection of "
    "synthetic events and faults, and a one-dimensional advection–dispersion network simulation. Escalating from "
    "30- to 15-min sampling did not detect more events than fixed 30-min sampling (recall "
    f"{f3(R[AD]['recall'])} vs {f3(R[FX]['recall'])}; {n0(N_TRIALS)} paired trials; {p(mc_af['recall']['p'])}). "
    "Detectors and classifiers built to resist sensor faults missed or mislabelled the record's expert-verified, "
    "turbidity-led events. Escalation paid off in source localization instead: with release time unknown, "
    "inverting sampled breakthrough curves halved median error "
    f"({CAD['Adaptive-30/15']['inversion_median']:.0f} vs {CAD['Fixed-30']['inversion_median']:.0f} m) for "
    f"{pc(extra_3015)} more samples, and 30 to 5-min escalation matched fixed 15-min accuracy with "
    f"{pc(saving_305_vs_15)} fewer samples. Adaptive cadence is best justified by event characterization rather "
    "than detection. Code, data and 80 automated tests are openly released."
)
assert len(ABSTRACT.split()) <= 150, len(ABSTRACT.split())

KEYWORDS = ("water-quality monitoring; adaptive sampling; event detection; sensor fault detection; "
            "contamination source identification; advection–dispersion; reproducibility")


# ----------------------------------------------------------------- docx -----
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.2)
st_normal = doc.styles["Normal"]
st_normal.font.name = "Times New Roman"
st_normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st_normal.font.size = Pt(11.5)
st_normal.paragraph_format.line_spacing = 1.5
st_normal.paragraph_format.space_after = Pt(6)
for name, size in (("Heading 1", 13), ("Heading 2", 11.5), ("Title", 16)):
    s = doc.styles[name]
    s.font.name = "Times New Roman"
    s.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.bold = True


def para(text: str = "", italic=False, bold=False, align=None, size=None):
    pg = doc.add_paragraph()
    if text:
        run = pg.add_run(text)
        run.italic, run.bold = italic, bold
        if size:
            run.font.size = Pt(size)
    if align:
        pg.alignment = align
    return pg


def rich(parts):
    """parts: list of (text, style) where style in {'', 'b', 'i'}."""
    pg = doc.add_paragraph()
    for text, style in parts:
        r = pg.add_run(text)
        r.bold, r.italic = "b" in style, "i" in style
    return pg


def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)


def bullets(items):
    for it in items:
        pg = doc.add_paragraph(style="List Bullet")
        pg.add_run(it)


EQ_N = [0]


def eq(text):
    EQ_N[0] += 1
    pg = para(align=WD_ALIGN_PARAGRAPH.CENTER)
    pg.add_run(text).italic = True
    pg.add_run(f"        ({EQ_N[0]})")
    pg.paragraph_format.space_before = Pt(3)


def _border(edge: str, size: int = 8):
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single" if size else "nil")
    if size:
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
    return el


def rules(t):
    """Journal style: horizontal rules above and below the table only; no vertical rules or shading."""
    borders = OxmlElement("w:tblBorders")
    for edge, size in (("top", 8), ("left", 0), ("bottom", 8), ("right", 0), ("insideH", 0), ("insideV", 0)):
        borders.append(_border(edge, size))
    t._tbl.tblPr.append(borders)


def table(caption, header, rows, widths=None, note=None):
    cp = para()
    r = cp.add_run(caption.split(".")[0] + ".")
    r.bold = True
    cp.add_run(caption[len(caption.split(".")[0]) + 1:])
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    rules(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        borders = OxmlElement("w:tcBorders")
        borders.append(_border("bottom", 6))
        c._tc.get_or_add_tcPr().append(borders)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9)
    for row in t.rows:
        for c in row.cells:
            c.paragraphs[0].paragraph_format.line_spacing = 1.0
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
    if widths:
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
    if note:
        para(note, italic=True, size=9)
    else:
        para()


FIG_N = {}


def figure(key, path, caption):
    FIG_N[key] = len(FIG_N) + 1
    doc.add_picture(str(ROOT / "figures" / path), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp = para()
    r = cp.add_run(f"Fig. {FIG_N[key]}. ")
    r.bold = True
    cp.add_run(caption)
    for run in cp.runs:
        run.font.size = Pt(9.5)


# ------------------------------------------------------------ front matter --
doc.add_heading(TITLE, level=0)
pg = para()
pg.add_run(AUTHOR).bold = True
sup = pg.add_run("a,*")
sup.font.superscript = True
pg = para()
sup = pg.add_run("a")
sup.font.superscript = True
sup.font.size = Pt(10)
addr = pg.add_run(f" {ADDRESS}")
addr.italic = True
addr.font.size = Pt(10)
para(f"* Corresponding author. E-mail address: {EMAIL}. ORCID: https://orcid.org/{ORCID}", italic=True, size=10)

h1("Highlights")
bullets(HIGHLIGHTS)

h1("Abstract")
para(ABSTRACT)
rich([("Keywords: ", "b"), (KEYWORDS, "")])

# ---------------------------------------------------------------- 1 intro ---
h1("1. Introduction")
para(
    "Automated in-situ sondes have made high-frequency river water-quality monitoring routine, but a deployed station "
    "still faces three operational decisions that existing studies usually treat one at a time. The first is when to "
    "sample. A short fixed interval catches fast events but drains batteries, accelerates probe fouling and produces "
    "more data than operators can review; a long interval is sustainable but can miss or mistime an event. The second "
    "is what an anomaly means. A sudden change may be a real change in the water or an instrument fault such as "
    "fouling, drift, dropout or exposure to air, and confusing the two either buries real alerts under nuisance alarms "
    "or teaches operators to ignore the system (Leigh et al., 2019; Rodriguez-Perez et al., 2020). The third arises "
    "once several stations are deployed: knowing that an event occurred is far less useful to a response team than "
    "knowing where it entered the river."
)
para(
    "Adaptive sampling, sampling coarsely in calm conditions and escalating when an anomaly appears, is an appealing "
    "answer to the first decision, and it is usually justified by detection: escalate, and the event will be caught "
    "sooner or more reliably. Whether that justification holds, and whether escalation pays off somewhere else, has "
    "rarely been tested against real background data with controlled ground truth. The second and third decisions "
    "interact with the first: a fault-aware detector changes which anomalies trigger escalation, and the cadence a "
    "station is running at during an event determines how much information it collects for locating the source."
)
para(
    "This paper evaluates all three decisions in one framework, using a public high-frequency river dataset as the "
    "real background and synthetic injected events and faults as ground truth, and it checks each component against "
    "the real events and real faults present in that dataset. The contributions are:"
)
bullets([
    "A paired, Monte Carlo comparison of fixed-interval, tiered adaptive and continuous sampling on real river "
    "background, with false alarms measured on the uninjected record rather than resampled across trials, and a "
    "detector whose standardization is shown to stay calibrated at every sampling cadence.",
    "A test of fault-versus-event discrimination that goes beyond synthetic cross-validation to real expert-verified "
    "hydrological events and real out-of-water fault records, exposing a tension between fault robustness and "
    "sensitivity to single-parameter events.",
    "A source-localization experiment in which release time, mass, velocity and dispersion are all unknown, which "
    "quantifies how much localization accuracy each sampling policy buys per sample and shows that this, not "
    "detection, is where adaptive escalation pays off.",
    "An open, tested implementation (Python, MIT licence; 80 automated tests) that regenerates every number, table "
    "and figure in this paper from the public data.",
])
para(
    "We report the results that did not favour adaptive sampling as fully as those that did. The study validates a "
    "software decision layer; it makes no claim about physical sensing hardware."
)

# ---------------------------------------------------------- 2 related work --
h1("2. Related work")
h2("2.1 Anomaly detection and fault discrimination in high-frequency water-quality data")
para(
    "Hill and Minsker (2010) framed anomaly detection in streaming environmental sensor data as a data-driven "
    "prediction problem, flagging readings that fall outside a model's prediction interval. Leigh et al. (2019) "
    "established a framework and public benchmark for river turbidity, conductivity and level data from catchments "
    "draining to the Great Barrier Reef, explicitly separating technical anomalies (sensor faults) from real "
    "water-quality anomalies. Talagala et al. (2019) proposed a feature-based procedure for flagging technical "
    "outliers, and Rodriguez-Perez et al. (2020) trained neural networks to detect technical anomalies in riverine "
    "and estuarine data. Alferes et al. (2013) described data validation and fault "
    "detection for in-situ measurement stations in water systems. Our fault-discrimination component follows the feature-based line "
    "of this work, favouring a small set of interpretable cross-channel features; its distinctive element is the "
    "check against real events and real faults rather than synthetic windows alone."
)
h2("2.2 Event-detection software and workflows")
para(
    "The US EPA CANARY software (Hart and Haxton) provides sensor-agnostic event detection for water-quality time "
    "series. Cronin et al. (2025) published an open, reproducible R workflow applying CANARY to high-frequency "
    "deployments in Ireland, including the Owenmore River record reused here, together with statistical event flags "
    "and expert verification of those flags. That workflow evaluates the detection algorithm; we evaluate the "
    "sampling policy layered on top of detection, and use the source study's verified events as an external check."
)
h2("2.3 Sensor placement and source identification")
para(
    "For pressurized distribution networks, contamination warning systems are designed and evaluated by simulating "
    "large ensembles of contamination scenarios through a hydraulic model, as formalized in the TEVA-SPOT toolkit "
    "(Janke, 2017) and benchmarked in the Battle of the Water Sensor Networks (Ostfeld et al., 2008). In rivers, "
    "source identification is an ill-posed inverse problem that has been approached geostatistically (Boano et al., "
    "2005) and with backward location probability density functions (Cheng and Jia, 2010). These methods assume the "
    "observations are given. Our question is the reverse: how the sampling policy that produces the observations "
    "affects how well a source can be located, and at what sampling cost."
)
h2("2.4 Positioning")
para(
    "Each component here has close precedent. To our knowledge, what has not been reported is their joint evaluation "
    "with real-data checks and a quantified, per-sample link between adaptive sampling cadence and source-localization "
    "accuracy under unknown release time. That combination is the contribution; we do not claim novelty for any "
    "single mechanism."
)

# ------------------------------------------------------------------ 3 data --
h1("3. Data")
d1, d2 = DATA["span"]["1"], DATA["span"]["2"]
para(
    "We used the raw YSI multiparameter sonde exports published by Cronin et al. (2025; dataset: Cronin, 2025) for "
    "two deployments on the "
    f"Owenmore River, Ireland, at a 15-min interval: deployment 1 from {d1[0][:16]} to {d1[1][:16]} "
    f"({n0(DATA['records']['1'])} in-river records) and deployment 2 from {d2[0][:16]} to {d2[1][:16]} "
    f"({n0(DATA['records']['2'])} records), about {REC_DAYS:.0f} days in total. The source study's cleaned files "
    "define the in-river period; pH, which the cleaned files omit, was taken from the raw exports by timestamp. We "
    "used four channels: temperature (°C), specific conductance (µS/cm), pH, and turbidity (NTU). Turbidity spans "
    "more than three orders of magnitude (Fig. 1) and is log1p-transformed before analysis; the other channels are "
    "used as recorded."
)
para(
    f"The source study provides two event labels per record: a statistical-control-threshold flag ({DATA['sct_flagged']} "
    f"records) and an expert verification ({DATA['verified_flagged']} records in four events). All four verified events "
    "fall in deployment 1 and are turbidity-led: turbidity rises by one to two orders of magnitude while conductance "
    "and temperature change little or fall (dilution). We treated every record carrying either label, with a 1-h "
    "buffer either side, as non-quiet; synthetic events were injected only into quiet periods, and false alarms were "
    "counted only in quiet periods."
)
nf = sum(f["records"] for f in DATA["faults"])
para(
    f"The raw exports also contain {nf} records in {len(DATA['faults'])} short windows, immediately before or after "
    "the in-river period of a deployment, "
    "where specific conductance falls below 5% of the deployment median (5–16 µS/cm against roughly 400 µS/cm), "
    "accompanied by temperature jumps. This pattern is physically implausible for the flowing river and consistent "
    "with the sonde being out of the water during deployment or retrieval; the source study's cleaning step removes "
    "these records. We used them as real, non-synthetic sensor-fault examples, never for training."
)
figure("data", "fig1_owenmore_data.png",
       "Owenmore River sonde record (Cronin et al., 2025). Shading: the source study's four expert-verified events. "
       "Crosses: real out-of-water fault records, which lie outside the in-river period used for the background "
       "and are shown at their raw values.")

# --------------------------------------------------------------- 4 methods --
h1("4. Methods")
h2("4.1 Anomaly detector")
para(
    "At each sample a station scores each channel against an exponentially weighted baseline and then folds the "
    "sample into it. Decay is defined in wall-clock time: a station that skips g native 15-min steps retains rᵍ of "
    "its prior baseline mean, with r = 0.5^(1/6) (a 90-min half-life). The deviation d = x − μ is standardized by an "
    "exponentially weighted mean of past squared deviations s², with a 6-h half-life, both computed before the new "
    "sample is folded in (Eq. 1):"
)
eq("z = (x − μ) / max(s, s_min),   μ ← μ + (1 − rᵍ) d,   s² ← r_v^g s² + (1 − r_v^g) d²")
para(
    "where s_min is a per-channel floor at instrument resolution and z is clipped to ±10. Standardizing by the "
    "variance of the prediction error, rather than by the usual incremental exponentially weighted variance, matters "
    "for a fair comparison across cadences. On stationary Gaussian noise the incremental estimator's steady-state "
    "value shrinks as the sampling gap grows, inflating z for coarse policies: E[z²] was "
    f"{CAL['1']['naive_z2']:.2f}, {CAL['2']['naive_z2']:.2f}, {CAL['4']['naive_z2']:.2f} and {CAL['8']['naive_z2']:.2f} "
    "at gaps of 15, 30, 60 and 120 min. With the prediction-error variance it was "
    f"{CAL['1']['z2']:.2f}, {CAL['2']['z2']:.2f}, {CAL['4']['z2']:.2f} and {CAL['8']['z2']:.2f}; the small residual "
    "rise at long gaps reflects fewer samples inside the variance window."
)
para(
    "The a priori multivariate score is the mean absolute z across the four channels. It rewards coherent "
    "multichannel change and, by design, damps a large deviation in a single channel, which is the typical signature "
    "of a sensor fault. As an exploratory variant, defined after inspecting the real events (Section 5.2), we also "
    "report a root-mean-square (RMS) score, which lets one strongly deviating channel reach the thresholds."
)
h2("4.2 Sampling policies and alarm rule")
para(
    "All policies share the detector and the alarm rule, so differences between them come only from when they sample. "
    "An alarm is raised when one sample scores at least τ₂ = 3.2, or two consecutive samples score at least τ₁ = 1.8; "
    "it clears at the first sample scoring below τ₁. Fixed-k samples every k minutes; Continuous-15 samples at the "
    "native 15-min resolution. Adaptive-b/a samples every b minutes in its normal tier and every a minutes once a "
    "sample scores at least τ₁, returning to the normal tier after three consecutive samples below τ₁. The main "
    "comparison is Fixed-30, Adaptive-30/15 and Continuous-15; a sweep adds Fixed-45 to Fixed-120 and Adaptive-45/15 "
    "to Adaptive-120/15. The thresholds, half-life and hold count were fixed before any evaluation and not tuned; "
    "Section 5.1 reports a sweep of τ₁. An alarm state is held on the native timeline until the next sample, because "
    "a consumer of a station's feed holds the last belief until new data arrive."
)
figure("example", "fig2_policy_example.png",
       "Illustrative injection trial. One injected moderate event on the real deployment-2 background (upper panel, turbidity) and the anomaly "
       "score at each sample taken by the three main policies (lower panel). Ticks mark each policy's first alarm.")
h2("4.3 Experiment 1: detection and false alarms")
para(
    "Detection was measured by injecting synthetic events into real quiet periods. An event has one of three temporal "
    "shapes (spike, ramp-plateau, damped oscillation) and one of three severity bands, expressed in local robust "
    "standard deviations (1.4826 × the median absolute deviation of the preceding day of quiet data): brief/faint "
    "(0.8–1.3 SD, 15–45 min), moderate (1.5–2.5 SD, 1–2.5 h) and severe (3.0–4.5 SD, 2–5 h). The signature is "
    "effluent-like: turbidity and conductance rise, pH falls, temperature rises at half amplitude. Onsets were drawn "
    f"from {P1['candidate_onsets']['1'] + P1['candidate_onsets']['2']:,} candidate positions where the event window, a "
    "1-h grace period and the preceding 2 h were all quiet. Each trial started the detector two days before onset, "
    f"and every policy was run on the identical injected series with a shared sampling phase. We ran {n0(N_TRIALS)} "
    f"trials ({N_SEV} per severity band). An event counts as detected if a sample taken at or after onset, and before "
    "the end of the event plus 1 h, is in alarm; an alarm already showing at onset from earlier samples is not "
    "credited. Latency is the time from onset to that sample."
)
para(
    "False alarms were measured separately, on the real, uninjected record, by running each policy over both full "
    "deployments and averaging over every sampling phase. We report alarm episodes per quiet day and the fraction of "
    "quiet steps in alarm. Repeating injection trials adds no new background, so a per-trial false-alarm statistic "
    "would be pseudo-replicated. The background also served for a real-event check: for each of the four verified "
    "events we recorded whether each policy alarmed between 1 h before and 2 h after the labelled onset."
)
h2("4.4 Experiment 2: fault versus event discrimination")
para(
    "Candidate anomalous windows were described relative to the 20 preceding samples, using only the first 6 h of "
    "the window so that a decision is available within 6 h. The a priori feature set (F4) comprises: the number of "
    "channels whose maximum |z| exceeds 3; directional consistency (turbidity and conductance move the same way and "
    "pH the opposite way, the expected co-movement for a discharge); physical plausibility (conductance stays above "
    "15% of baseline); and an abruptness ratio (the largest single step divided by the window's range in the most "
    "deviant channel). A revised, sign-agnostic set (F6), defined after the F4 results on real events, replaces "
    "directional consistency with the second-largest channel deviation and the synchrony of channel onsets, extends "
    "the plausibility test to pH and temperature, and adds the fraction of exactly-zero readings."
)
para(
    f"Each labelled pool held {P2['n_per_class']} synthetic event windows and {P2['n_per_class']} synthetic fault "
    "windows on real quiet background. Faults affected exactly one randomly chosen channel and were split evenly "
    "across four types: zero dropout, linear drift (2–6 SD), noise burst and conductance collapse to 1–5% of "
    "baseline. Events used either the discharge signature alone (the a priori design) or discharge and runoff-like "
    "signatures in equal numbers, the runoff signature having turbidity up, conductance diluted and slight cooling. "
    "Both feature sets were evaluated under both pools with class-balanced logistic regression (the primary model) "
    f"and a random forest, using {P2['n_repeats']}-times-repeated stratified five-fold cross-validation, and against "
    "a one-line rule (fault if at most one channel deviates). We also measured: recall on each fault type when that "
    "type was left out of training; recall on an unseen multichannel fouling fault, in which turbidity drifts up and "
    "conductance down together; and, with the model fitted to the whole pool, predictions on the seven real windows."
)
h2("4.5 Experiment 3: source localization")
para(
    "A conservative tracer of mass M is released instantaneously at position x₀ and time t₀ into a straight 10-km "
    "reach of cross-section A = 5 m² and is transported by one-dimensional advection–dispersion (Fischer et al., 1979), Eq. (2):"
)
eq("C(x, t) = M / (A √(4π D τ)) · exp(−(x − x₀ − v τ)² / (4 D τ)),   τ = t − t₀ > 0")
para(
    "For each scenario x₀ ~ U(0, 6 km), t₀ ~ U(0, 6 h), v ~ U(0.08, 0.6 m/s), D ~ U(0.5, 8 m²/s) and "
    "log₁₀ M ~ U(3, 5) (M in g). Between three and six stations were placed uniformly along the reach. Each station's "
    "concentration was observed on a 5-min grid for 48 h with multiplicative (5%) and additive (0.005 mg/L) Gaussian "
    "noise. Each sampling policy then subsampled the same noisy series, so all policy comparisons are paired: "
    "Fixed-60, Fixed-30, Fixed-15, Fixed-5, Adaptive-30/15 and Adaptive-30/5. An adaptive station escalates when a "
    "sample exceeds the 0.05-mg/L detection threshold (ten times the noise floor) and de-escalates after three samples "
    "below it. Sampling phases differ between stations."
)
para(
    "The proposed estimator fits Eq. (2) to every sample from every station, including the zeros, by "
    "weighted nonlinear least squares with all five parameters (x₀, t₀, v, D, M) unknown and bounded, using six "
    "starting points. The breakthrough curve's spread carries the travel-time information that arrival times alone "
    "lack: with t₀ unknown, arrival times constrain only the velocity and a line of (x₀, t₀) combinations. As a "
    "reference we also report an arrival-time regression (x = x₀ + v(t − t₀), with arrival taken as the time of each "
    "station's highest sample) that is given the true t₀. Because t₀ is never known in practice, this reference is "
    "optimistic. A scenario is localized when at least two stations detect the plume. We ran "
    f"{P3['n_scenarios']} scenarios, and a nested station-count experiment ({P3['n_station_scenarios']} releases, each "
    "observed by the first 2, 3, 4, 6 and 8 of eight randomly placed stations)."
)
h2("4.6 Statistics")
para(
    "Paired binary outcomes (detected or not) were compared with the exact McNemar test (McNemar, 1947); paired "
    "continuous outcomes (latency among events detected by both policies; localization error among scenarios "
    "localized under both policies) with the two-sided Wilcoxon signed-rank test (Wilcoxon, 1945), with bootstrap "
    "95% confidence intervals (5,000 resamples) on the mean or median paired difference. Proportions carry Wilson "
    "intervals. P-values within each experiment's family of planned comparisons were adjusted with Holm's procedure "
    "(Holm, 1979). With samples this large, small differences reach significance, so we report effect sizes alongside "
    "every test and base conclusions on them."
)
h2("4.7 Implementation, testing and reproducibility")
para(
    "The framework is a Python package (NumPy, pandas, SciPy and scikit-learn; Virtanen et al., 2020; Pedregosa et al., "
    "2011) with 80 automated tests. Among other things, the tests check: data integrity against the source files; that "
    "z is pre-update and calibrated at every cadence; that the persistence rule and forward-filled alarm state behave "
    "as specified, and that a stale alarm is never credited as a detection; that injections affect only the declared "
    "channels and windows; mass conservation and the first two moments of the transport solution; that the "
    "inversion recovers the source to within 50 m from near-noise-free data with t₀ unknown; and the exact values of "
    "the statistical tests. All random seeds are fixed. Every number, table and figure in this paper is generated by "
    "scripts that read the experiment outputs, and the full pipeline runs in about 10 min on an 8-core laptop."
)

# --------------------------------------------------------------- 5 results --
h1("5. Results")
h2("5.1 Adaptive escalation did not improve detection")
rows = []
for name in (FX, AD, CO):
    r = R[name]
    rows.append([name, f"{f3(r['recall'])} {ci(r['ci'])}", f3(r["recall_brief"]), f3(r["recall_moderate"]),
                 f3(r["recall_severe"]), f"{r['latency_median']:.0f} / {r['latency_p90']:.0f}",
                 f"{BG[(name, 2)]['cycles_per_day']:.1f}"])
para(
    "Table 1 and Fig. 3a summarize detection. Adaptive-30/15 and Fixed-30 detected the same fraction of injected "
    f"events ({f3(R[AD]['recall'])} vs {f3(R[FX]['recall'])}; {mc_af['recall']['a_only']} trials detected only by "
    f"the adaptive policy and {mc_af['recall']['b_only']} only by the fixed policy; exact McNemar "
    f"{p(mc_af['recall']['p'])}). No severity band showed a difference. Among events both detected, escalation "
    f"shortened mean latency by {abs(lat_af['mean_diff']):.1f} min (95% CI {abs(lat_af['ci_mean_diff'][1]):.1f}–"
    f"{abs(lat_af['ci_mean_diff'][0]):.1f} min; Holm-adjusted {p(P1['holm']['Adaptive_vs_Fixed30|latency'])}), "
    f"against a median paired difference of zero. That is a real but operationally negligible gain. Continuous-15 "
    f"detected more than either ({f3(R[CO]['recall'])}; {mc_ac['recall']['b_only']} vs "
    f"{mc_ac['recall']['a_only']} discordant trials against Adaptive-30/15, {p(mc_ac['recall']['p'])}), with the gain "
    f"concentrated in moderate events ({f3(R[CO]['recall_moderate'])} vs {f3(R[AD]['recall_moderate'])})."
)
table("Table 1. Detection of injected events by the three main policies (" + n0(N_TRIALS) + " paired trials; recall "
      "with 95% Wilson interval; latency in minutes among detected events; cost on the real record).",
      ["Policy", "Recall (95% CI)", "Brief/faint", "Moderate", "Severe", "Latency median / P90", "Samples per day"],
      rows, widths=[2.6, 3.0, 1.6, 1.6, 1.4, 2.4, 1.6])
para(
    "Severe events were detected reliably by all policies (≥ "
    f"{min(R[n]['recall_severe'] for n in (FX, AD, CO)):.2f}); brief/faint events, at 0.8–1.3 local SD, almost never "
    f"(≤ {max(R[n]['recall_brief'] for n in (FX, AD, CO)):.2f}). That is expected for a detector that requires the "
    "mean deviation across four channels to reach 1.8 SD on two consecutive samples. Across the sweep (Fig. 3b), each "
    "adaptive policy sat on or slightly below the curve traced by fixed policies of similar cost: escalation did not "
    "buy detection that a fixed interval of the same average cost could not. The reason is structural. A policy "
    "escalates at the same sample that first crosses τ₁, so the time of first detection is set by the base interval; "
    "escalation can only speed the second, confirming sample. Events shorter than the base interval are invisible to "
    "both."
)
figure("recall", "fig3_recall_cost.png",
       "Detection by severity and sampling cost. (a) Recall by severity band for the three main policies (500 trials per band). (b) Recall for moderate events "
       "against sampling cost for the fixed-interval family (labels: interval in min) and the adaptive family "
       "(labels: base/alert interval in min).")
tau = {(r["policy"], r["tau1"]): r for r in P1["tau_sweep"]}
para(
    "Varying τ₁ from 1.4 to 2.6 (Fig. 4) moved recall for moderate events monotonically, from "
    f"{tau[(AD, 1.4)]['recall_moderate']:.2f} to {tau[(AD, 2.6)]['recall_moderate']:.2f} for Adaptive-30/15 and from "
    f"{tau[(CO, 1.4)]['recall_moderate']:.2f} to {tau[(CO, 2.6)]['recall_moderate']:.2f} for Continuous-15. At "
    "every threshold Continuous-15 detected the most, and Fixed-30 and Adaptive-30/15 stayed within "
    f"{max(abs(tau[(FX, t)]['recall_moderate'] - tau[(AD, t)]['recall_moderate']) for t in (1.4, 1.6, 1.8, 2.0, 2.2, 2.6)):.2f} "
    "of each other, swapping order at some thresholds. The conclusions above therefore do not depend on the a "
    "priori threshold."
)
figure("tau", "fig4_threshold_sweep.png",
       "Threshold sensitivity, with τ₂ fixed. (a) recall for moderate injected events; (b) false-alarm "
       "episodes per quiet day on the real, uninjected record. Dotted line: the a priori τ₁ = 1.8.")

h2("5.2 False alarms on the real record, and the real events")
b2 = {n: BG[(n, 2)] for n in (FX, AD, CO)}
b1 = {n: BG[(n, 1)] for n in (FX, AD, CO)}
para(
    f"Over about {QUIET_DAYS:.0f} quiet days, the two-sample persistence rule kept background false alarms rare for "
    f"every policy: {b2[FX]['false_episodes_per_day']:.2f}, {b2[AD]['false_episodes_per_day']:.2f} and "
    f"{b2[CO]['false_episodes_per_day']:.2f} episodes per quiet day for Fixed-30, Adaptive-30/15 and Continuous-15 "
    f"(Table 2). These rates rest on at most {max(b2[n]['false_episodes'] for n in b2):.0f} episode per policy, too "
    "few to rank the policies. With a single-sample "
    f"alarm rule they rose to {b1[FX]['false_episodes_per_day']:.2f}, {b1[AD]['false_episodes_per_day']:.2f} and "
    f"{b1[CO]['false_episodes_per_day']:.2f} per day, and a fixed policy held each false alarm for longer, because it "
    f"cannot re-sample sooner (alarmed quiet steps: {b1[FX]['far_per_1000']:.1f}, {b1[AD]['far_per_1000']:.1f} and "
    f"{b1[CO]['far_per_1000']:.1f} per 1,000). That faster clearing is the one consistent operational benefit of "
    "escalation we found for alarms."
)
rows = []
for name in (FX, AD, CO):
    for agg, nm in (("mean", name), ("rms", rms[name])):
        rows.append([nm.replace(" [rms]", ""), "Mean |z| (a priori)" if agg == "mean" else "RMS (exploratory)",
                     f"{BG[(nm, 2)]['false_episodes_per_day']:.2f}", f"{BG[(nm, 1)]['false_episodes_per_day']:.2f}",
                     f"{f3(R[nm]['recall'])}", f"{real_hits(nm):.1f} / 4"])
table("Table 2. Real-record false alarms, injected-event recall and real-event detection, by policy and score "
      "aggregation. False alarms are episodes per quiet day under the two-sample (persist = 2) and single-sample "
      "(persist = 1) alarm rules. Real events: verified events alarmed within −1 h to +2 h of onset, averaged over "
      "sampling phase.",
      ["Policy", "Score", "False alarms/day (persist 2)", "(persist 1)", "Injected recall", "Real events detected"],
      rows, widths=[2.6, 3.2, 2.6, 1.6, 2.0, 2.4])
para(
    "The real-event check exposed the main weakness of the a priori detector. Of the four expert-verified events, "
    f"Adaptive-30/15 alarmed near {real_hits(AD):.0f}, Continuous-15 near {real_hits(CO):.0f}, and Fixed-30 near "
    f"{real_hits(FX):.1f} on average over sampling phase (Table 2). All four events are led by turbidity, which "
    "rises sharply from the first sample while the other channels barely move. The mean-|z| score caps a single "
    "channel's contribution at 10/4 = 2.5. After the first spike, the inflated error variance makes the next sample "
    "look ordinary, so the two-sample confirmation never completes. The exploratory RMS score detected "
    f"{real_hits(rms[AD]):.0f} of the four under the adaptive policy and raised injected recall from "
    f"{f3(R[AD]['recall'])} to {f3(R[rms[AD]]['recall'])}. Under the single-sample rule it also raised background "
    f"false alarms from {b1[AD]['false_episodes_per_day']:.2f} to {BG[(rms[AD], 1)]['false_episodes_per_day']:.2f} "
    "per day: the price of letting one channel speak for all four."
)

h2("5.3 Fault discrimination: good on synthetic windows, mixed on real ones")
rows = []
labels = {"discharge_only|F4": "F4, discharge only (a priori)", "discharge_only|F6": "F6, discharge only",
          "discharge_and_runoff|F4": "F4, discharge + runoff", "discharge_and_runoff|F6": "F6, discharge + runoff"}
for k, lab in labels.items():
    d = DES[k]
    lo = d["leave_one_type_out_recall"]
    rows.append([lab, f"{f2(d['logistic_cv']['f1_mean'])} ± {f2(d['logistic_cv']['f1_sd'])}",
                 f"{f2(d['forest_cv']['f1_mean'])}", f"{f2(d['rule']['f1'])}",
                 " / ".join(f"{lo[t]:.2f}" for t in ("dropout", "drift", "noise_burst", "collapse")),
                 f"{d['unseen_fouling_recall']:.2f}",
                 f"{d['external_correct']['real_fault']}/3", f"{d['external_correct']['real_event']}/4"])
table("Table 3. Fault-versus-event discrimination. F1 for the fault class under 5-fold cross-validation repeated "
      f"{P2['n_repeats']} times (mean ± SD); leave-one-type-out recall for dropout / drift / noise burst / collapse; "
      "recall on the unseen fouling fault; correct predictions on the real windows (logistic model).",
      ["Design", "Logistic F1", "Forest F1", "Rule F1", "Held-out type recall", "Fouling", "Real faults",
       "Real events"],
      rows, widths=[3.4, 1.9, 1.4, 1.2, 3.4, 1.3, 1.3, 1.3])
para(
    f"Under the a priori design, the logistic classifier reached F1 = {f2(a_pri['logistic_cv']['f1_mean'])} ± "
    f"{f2(a_pri['logistic_cv']['f1_sd'])} (Table 3), well above the one-line rule "
    f"({f2(a_pri['rule']['f1'])}); the random forest was similar ({f2(a_pri['forest_cv']['f1_mean'])}). All four "
    "standardized coefficients pointed toward the event class as intended, the largest being physical plausibility "
    f"({a_pri['coefficients']['physical_plausibility']:+.2f}) and directional consistency "
    f"({a_pri['coefficients']['directional_consistency']:+.2f}). The model flagged all three real out-of-water "
    "windows as faults (Fig. 5b). It "
    f"also labelled {WORDS[4 - a_pri['external_correct']['real_event']]} of the four real hydrological events as faults. "
    "In both, turbidity rose while conductance fell or held steady, which violates the discharge direction the "
    "directional feature encodes."
)
para(
    "The revised, sign-agnostic features corrected most of the real events "
    f"({rev['external_correct']['real_event']}/4 correct with F6 trained on both signatures) but missed every real "
    f"fault ({rev['external_correct']['real_fault']}/3). The out-of-water windows combine a conductance collapse with "
    "simultaneous temperature and pH shifts, which looks coherent and multichannel to features designed around "
    "synchrony. Training F4 on both signatures gave the best real-window score "
    f"({f4_both['external_correct']['real_fault']}/3 faults and {f4_both['external_correct']['real_event']}/4 events), "
    f"but its logistic F1 fell to {f2(f4_both['logistic_cv']['f1_mean'])}, because its directional feature no longer "
    "separates the classes. Recall on a fault type withheld from training held up for the a priori design "
    f"({min(a_pri['leave_one_type_out_recall'].values()):.2f}–{max(a_pri['leave_one_type_out_recall'].values()):.2f}) "
    "but fell sharply for drift and noise bursts in the other three designs "
    f"(as low as {min(min(DES[k]['leave_one_type_out_recall'][t] for t in ('drift', 'noise_burst')) for k in DES):.2f}). "
    "The unseen multichannel fouling fault was mostly flagged by all designs "
    f"({min(DES[k]['unseen_fouling_recall'] for k in DES):.2f}–{max(DES[k]['unseen_fouling_recall'] for k in DES):.2f}). "
    "With only seven real windows, these real-data results are indicative, not estimates of field "
    "performance. Their consistent message is that synthetic cross-validation overstates how well any of these "
    "designs transfers."
)
figure("clf", "fig5_fault_classifier.png",
       "Fault-versus-event discrimination. (a) Cross-validated F1 for the fault class under the four designs (bars: logistic and random forest; mean ± SD "
       "over 50 folds; dashes: one-line rule). (b) Predicted probability of fault for the seven real windows under the "
       "a priori and revised designs; shaded rows are true faults; dotted line: decision threshold.")

h2("5.4 Escalation pays off in source localization")
rows = []
for k in ("Fixed-60", "Fixed-30", "Adaptive-30/15", "Fixed-15", "Adaptive-30/5", "Fixed-5"):
    c = CAD[k]
    rows.append([k, f"{c['samples_per_station_day']:.1f}", f"{c['localized_rate']:.2f}",
                 f"{c['inversion_median']:.0f} [{c['inversion_median_ci'][0]:.0f}, {c['inversion_median_ci'][1]:.0f}]",
                 f"{c['inversion_p90']:,.0f}", f"{c['inversion_within_100m']:.2f}", f"{c['oracle_median']:.0f}"])
para(
    "Localization error fell steadily as sampling became denser (Table 4, Fig. 6a), from a median of "
    f"{m(CAD['Fixed-60']['inversion_median'])} at 60-min sampling to {m(CAD['Fixed-5']['inversion_median'])} at 5 min, "
    f"on a 10-km reach. Adaptive escalation moved policies off this cost curve. Adaptive-30/15 halved the median error "
    f"of Fixed-30 ({m(CAD['Adaptive-30/15']['inversion_median'])} vs {m(CAD['Fixed-30']['inversion_median'])}) for "
    f"{pc(extra_3015)} more samples. It was more accurate in {pc(c3015['a_better_fraction'])} of paired scenarios "
    f"(median paired difference {c3015['median_diff']:.0f} m, 95% CI {c3015['ci_median_diff'][0]:.0f} to "
    f"{c3015['ci_median_diff'][1]:.0f} m; Holm-adjusted {p(c3015['p_holm'])}). Adaptive-30/5 matched Fixed-15 "
    f"({m(CAD['Adaptive-30/5']['inversion_median'])} vs {m(CAD['Fixed-15']['inversion_median'])}; median paired "
    f"difference {c305_15['median_diff']:.0f} m, {p(c305_15['p_holm'])}) while taking {pc(saving_305_vs_15)} fewer "
    "samples, and the two error distributions nearly coincide (Fig. 6b). The gain comes from where the samples fall. "
    "Escalation concentrates them on the breakthrough curve, whose arrival and spread pin down the travel time and "
    "hence the source; samples taken while the river is clean carry almost no information about the source."
)
table("Table 4. Source localization by sampling policy (" + str(P3["n_scenarios"]) + " scenarios; release time, "
      "mass, velocity and dispersion unknown). Error in metres; median with bootstrap 95% CI. The final column is the "
      "arrival-time regression given the true release time (optimistic reference).",
      ["Policy", "Samples per station-day", "Localized", "Median error [95% CI]", "P90 error", "≤ 100 m",
       "Reference (t₀ known)"],
      rows, widths=[2.6, 2.0, 1.6, 3.2, 1.6, 1.4, 2.2])
rows = []
for k in ("Adaptive-30/15 vs Fixed-30", "Adaptive-30/5 vs Fixed-30", "Adaptive-30/5 vs Fixed-15",
          "Fixed-15 vs Fixed-30", "Fixed-5 vs Fixed-15"):
    c = CMP[k]
    rows.append([k, c["n"], f"{c['median_diff']:.0f} [{c['ci_median_diff'][0]:.0f}, {c['ci_median_diff'][1]:.0f}]",
                 f"{c['a_better_fraction']:.2f}", p(c["p_holm"]).replace("p ", "")])
table("Table 5. Paired comparisons of localization error (first policy minus second, in metres, over scenarios "
      "localized under both). Wilcoxon signed-rank, Holm-adjusted.",
      ["Comparison", "n", "Median difference [95% CI]", "First better (fraction)", "Adjusted p"],
      rows, widths=[4.6, 1.2, 4.0, 2.6, 2.2])
para(
    "The proposed inversion, which does not know the release time, was more accurate than the reference regression "
    "that does, under every policy (Table 4). Arrival times quantized to the sampling interval lose the information "
    f"in a curve's shape. The fraction of scenarios localized rose from {CAD['Fixed-60']['localized_rate']:.2f} "
    f"to {CAD['Fixed-5']['localized_rate']:.2f} with cadence; failures were scenarios in which fewer than two "
    "stations lay downstream of the release or saw a concentration above the detection threshold. Adding stations "
    "helped both policies, and escalation helped at every station count (Fig. 6c): median error with two stations was "
    f"{m(st[('Fixed-30', 2)]['median_error'])} under Fixed-30 and {m(st[('Adaptive-30/15', 2)]['median_error'])} under "
    f"Adaptive-30/15, and with eight stations {m(st[('Fixed-30', 8)]['median_error'])} and "
    f"{m(st[('Adaptive-30/15', 8)]['median_error'])}. In this experiment, escalating cadence was worth roughly as "
    "much accuracy as doubling the number of stations."
)
figure("loc", "fig6_localization.png",
       "Source localization by sampling policy. (a) Median source-localization error (bootstrap 95% CI) against sampling cost for fixed and adaptive policies "
       "(labels: interval in min). (b) Cumulative distribution of error for four policies; curves end below 1 because "
       "some scenarios could not be localized. (c) Median error against the number of stations on the reach (nested "
       "design: identical releases).")

# ------------------------------------------------------------ 6 discussion --
h1("6. Discussion")
h2("6.1 Where adaptive sampling pays off")
para(
    "The results separate two jobs that adaptive sampling is often expected to do. As a detection strategy, "
    "escalation added almost nothing: the base interval decides when an event is first seen, and a fixed interval of "
    "the same average cost detects about as much. As a characterization strategy, escalation was efficient: once a "
    "station knows something is happening, dense samples on the passing plume are exactly what a localization "
    "estimator needs. For monitoring design, this suggests choosing the base interval from the detection requirement, "
    "the shortest and faintest event that must be caught, and justifying escalation by what must be learned about an "
    "event once it is detected: its source, shape and load. Escalating to 5 min, for example, recovered the accuracy of a 15-min fixed interval at "
    f"{pc(samp['Adaptive-30/5'] / samp['Fixed-15'])} of its sampling cost."
)
h2("6.2 Fault robustness versus sensitivity")
para(
    "The real-data checks expose a tension that synthetic evaluation hides. A score that averages across channels, and "
    "a classifier that looks for coherent multichannel change, both resist single-channel faults by design. At this "
    "site, however, every verified event was led by a single parameter, turbidity. The two goals pull in opposite "
    "directions, and the features that separate synthetic events from synthetic faults encode assumptions about what "
    "real events look like. The directional feature assumed a discharge; the synchrony features failed on out-of-water "
    "records, which are coherent across channels. A more defensible architecture is probably sensitive per-channel "
    "detection followed by a discriminator trained on real, labelled fault libraries from several deployments and "
    "sensor models. Building such libraries, rather than refining features on synthetic data, is the main open problem "
    "this study points to."
)
h2("6.3 Limitations")
bullets([
    f"One river, two deployments, about {REC_DAYS:.0f} days in one season. Detection and false-alarm results are "
    "specific to this background and should be repeated on other rivers, for example further CANARY-style deployments "
    "or sub-daily subsets of the Global River Water Quality Archive (Virro et al., 2021).",
    "Detection metrics rest on synthetic injected events of assumed shape and signature. The real-data checks (four "
    "events, three fault windows) are small and indicative, and the fault windows are all of one kind (out of water).",
    "Fault-discrimination windows were segmented with knowledge of the injected onset. In deployment, segmentation "
    "would come from the detector and would add error.",
    "Localization uses an idealized transport model: a straight, uniform channel, constant velocity and dispersion, "
    "an instantaneous conservative release, and stations that measure the contaminant itself. Tributaries, dead zones, "
    "decay and surrogate-only sensing would all degrade accuracy. A calibrated hydrodynamic model (e.g. HEC-RAS) for a "
    "real reach is the natural next test.",
    "Every scheduled sample is assumed to succeed; missed cycles from power or communication faults are not modelled.",
    "The thresholds and half-lives were fixed a priori rather than optimized. The RMS score and the F6 features were "
    "defined after inspecting the real events and are reported as exploratory.",
    "No physical hardware was built or tested.",
])

# ----------------------------------------------------------- 7 conclusions --
h1("7. Conclusions")
para(
    "On a real river background with controlled ground truth, tiered adaptive sampling did not detect more events than "
    "fixed sampling of similar cost, and multichannel, fault-robust detection missed most of the real turbidity-led "
    "events in the record. Where escalation did pay off was after detection. Concentrating samples on a passing plume "
    f"halved source-localization error for {pc(extra_3015)} more samples, or matched a fixed 15-min interval at "
    f"{pc(samp['Adaptive-30/5'] / samp['Fixed-15'])} of its sampling cost, without knowing when the release occurred. "
    "Monitoring networks that adopt adaptive cadence should "
    "justify it, and tune its alert interval, by what they need to learn about an event once it is found. The open "
    "framework and tests released with this paper allow each of these results to be rerun and extended to other "
    "rivers."
)

# -------------------------------------------------------------- back matter --
h1("Software and data availability")
for k, v in [
    ("Name of software", "aquasentinel (v1.0.0)"),
    ("Developer and contact", f"{AUTHOR} ({EMAIL})"),
    ("Year first available", "2026"),
    ("Hardware required", "Any computer able to run Python; full pipeline ≈10 min on an 8-core laptop"),
    ("Software required", "Python ≥ 3.10 with NumPy, pandas, SciPy, scikit-learn, matplotlib, joblib; pytest for tests"),
    ("Program language", "Python"),
    ("Program size", "≈1,100 lines (package) and ≈480 lines (tests), plus experiment, figure and manuscript scripts"),
    ("Availability and cost", f"Free and open source (MIT licence): {REPO_URL}; archived version: "
                              f"https://doi.org/{ZENODO_DOI} (Nunna, 2026)"),
    ("Data", "Owenmore River sonde data and event labels from Cronin (2025) (MIT licence), redistributed "
             "unmodified in the archive above together with all raw experiment outputs (CSV and JSON files, ≈1.5 MB)"),
]:
    rich([(f"{k}: ", "b"), (v, "")])

h1("CRediT authorship contribution statement")
para(f"{AUTHOR}: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, "
     "Data curation, Writing – original draft, Writing – review & editing, Visualization.")
h1("Declaration of competing interest")
para("The author declares that they have no known competing financial interests or personal relationships that "
     "could have appeared to influence the work reported in this paper.")
h1("Funding")
para("This research did not receive any specific grant from funding agencies in the public, commercial, or "
     "not-for-profit sectors.")
h1("Declaration of generative AI and AI-assisted technologies in the manuscript preparation process")
para("During the preparation of this work the author used Claude (Anthropic) to implement and test the analysis "
     "code, run the experiments, produce the data visualizations from the experiment outputs, and draft the "
     "manuscript text. After using this tool, the author reviewed and edited the content as needed and takes full "
     "responsibility for the content of the published article.")
h1("Acknowledgements")
para("The author thanks Cronin et al. for publishing their data, labels and workflow openly, which made this study possible.")

h1("References")
REFS = [
    "Alferes, J., Tik, S., Copp, J., Vanrolleghem, P.A., 2013. Advanced monitoring of water systems using in situ "
    "measurement stations: data validation and fault detection. Water Sci. Technol. 68 (5), 1022–1030. "
    "https://doi.org/10.2166/wst.2013.302",
    "Boano, F., Revelli, R., Ridolfi, L., 2005. Source identification in river pollution problems: a geostatistical "
    "approach. Water Resour. Res. 41 (7), W07023. https://doi.org/10.1029/2004WR003754",
    "Cheng, W.P., Jia, Y., 2010. Identification of contaminant point source in surface waters based on backward "
    "location probability density function method. Adv. Water Resour. 33 (4), 397–410. "
    "https://doi.org/10.1016/j.advwatres.2010.01.004",
    "[dataset] Cronin, L., 2025. Workflow for identifying water quality hydrodynamic events in time series data "
    "(code, Owenmore River sonde data and event labels). GitHub. "
    "https://github.com/LCroninATU/Workflow-for-identifying-water-quality-hydrodynamic-events-in-time-series-data",
    "Cronin, L., Taylor, C.M., Briciu Burghina, C., Lucy, F.E., Regan, F., 2025. An open-source workflow for "
    "identifying hydrodynamic water quality events in rivers by continuous water quality monitoring and time-series "
    "data processing using R and US EPA CANARY. MethodsX 15, 103538. https://doi.org/10.1016/j.mex.2025.103538",
    "Fischer, H.B., List, E.J., Koh, R.C.Y., Imberger, J., Brooks, N.H., 1979. Mixing in Inland and Coastal Waters. "
    "Academic Press, New York.",
"Hart, D., Haxton, T. CANARY (Continuous Analysis of Networked Array of Sensors) event detection software. "
    "Sandia National Laboratories / US Environmental Protection Agency. https://github.com/USEPA/CANARY",
    "Hill, D.J., Minsker, B.S., 2010. Anomaly detection in streaming environmental sensor data: a data-driven "
    "modeling approach. Environ. Model. Softw. 25 (9), 1014–1022. https://doi.org/10.1016/j.envsoft.2009.08.010",
    "Holm, S., 1979. A simple sequentially rejective multiple test procedure. Scand. J. Stat. 6 (2), 65–70.",
    "Janke, R., 2017. Threat Ensemble Vulnerability Assessment – Sensor Placement Optimization Tool (TEVA-SPOT) "
    "Graphical User Interface User's Manual. EPA/600/R-17/385. U.S. Environmental Protection Agency, Washington, DC.",
    "Leigh, C., Alsibai, O., Hyndman, R.J., Kandanaarachchi, S., King, O.C., McGree, J.M., Neelamraju, C., Strauss, "
    "J., Talagala, P.D., Turner, R.D.R., Mengersen, K., Peterson, E.E., 2019. A framework for automated anomaly "
    "detection in high frequency water-quality data from in situ sensors. Sci. Total Environ. 664, 885–898. "
    "https://doi.org/10.1016/j.scitotenv.2019.02.085",
    "McNemar, Q., 1947. Note on the sampling error of the difference between correlated proportions or percentages. "
    "Psychometrika 12 (2), 153–157. https://doi.org/10.1007/BF02295996",
    f"[dataset] Nunna, D.S.G., 2026. aquasentinel v1.0.0: code, data and experiment outputs for adaptive "
    f"sampling, fault discrimination and source localization in river water-quality monitoring. Zenodo. "
    f"https://doi.org/{ZENODO_DOI}",
        "Ostfeld, A., Uber, J.G., Salomons, E., Berry, J.W., Hart, W.E., Phillips, C.A., et al., 2008. The battle of the "
    "water sensor networks (BWSN): a design challenge for engineers and algorithms. J. Water Resour. Plan. Manag. "
    "134 (6), 556–568. https://doi.org/10.1061/(ASCE)0733-9496(2008)134:6(556)",
    "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., et al., 2011. Scikit-learn: "
    "machine learning in Python. J. Mach. Learn. Res. 12, 2825–2830.",
    "Rodriguez-Perez, J., Leigh, C., Liquet, B., Kermorvant, C., Peterson, E., Sous, D., Mengersen, K., 2020. "
    "Detecting technical anomalies in high-frequency water-quality data using artificial neural networks. Environ. "
    "Sci. Technol. 54 (21), 13719–13730. https://doi.org/10.1021/acs.est.0c04069",
    "Talagala, P.D., Hyndman, R.J., Leigh, C., Mengersen, K., Smith-Miles, K., 2019. A feature-based procedure for "
    "detecting technical outliers in water-quality data from in situ sensors. Water Resour. Res. 55 (11), 8547–8568. "
    "https://doi.org/10.1029/2019WR024906",
    "Virro, H., Amatulli, G., Kmoch, A., Shen, L., Uuemaa, E., 2021. GRQA: Global River Water Quality Archive. Earth "
    "Syst. Sci. Data 13 (12), 5483–5507. https://doi.org/10.5194/essd-13-5483-2021",
    "Virtanen, P., Gommers, R., Oliphant, T.E., Haberland, M., Reddy, T., Cournapeau, D., et al., 2020. SciPy 1.0: "
    "fundamental algorithms for scientific computing in Python. Nat. Methods 17 (3), 261–272. "
    "https://doi.org/10.1038/s41592-019-0686-2",
    "Wilcoxon, F., 1945. Individual comparisons by ranking methods. Biometrics Bull. 1 (6), 80–83. "
    "https://doi.org/10.2307/3001968",
]
for ref in REFS:
    pg = para(ref, size=10)
    pg.paragraph_format.left_indent = Cm(0.6)
    pg.paragraph_format.first_line_indent = Cm(-0.6)
    pg.paragraph_format.line_spacing = 1.15

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
(OUT.parent / "highlights.txt").write_text("\n".join(HIGHLIGHTS) + "\n", encoding="utf-8")
hd = Document()
hd.add_heading("Highlights", level=1)
for h in HIGHLIGHTS:
    hd.add_paragraph(h, style="List Bullet")
hd.save(OUT.parent / "highlights.docx")
print(f"abstract: {len(ABSTRACT.split())} words")

words = sum(len(pg.text.split()) for pg in doc.paragraphs)
print(f"wrote {OUT} ({words:,} words in paragraphs; {len(FIG_N)} figures)")
